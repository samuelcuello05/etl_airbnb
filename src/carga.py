"""
Fase de CARGA del proceso ETL.

Contiene la clase ``Carga``, que toma los DataFrames limpios producidos por
``Transformacion`` y:

1. los inserta en una base de datos SQLite nueva (data/processed/airbnb_bogota.db),
2. los exporta a archivos XLSX,
3. verifica que la cantidad de registros cargados coincida con la de origen,
4. registra en el log cada evento del proceso.
"""

import re
import sqlite3
import time
from pathlib import Path

import pandas as pd
import xlsxwriter
from openpyxl import load_workbook

from config import (MAX_FILAS_EXCEL, RUTA_SQLITE, RUTA_XLSX_RESUMEN,
                    RUTA_XLSX_REVIEWS)
from logger import obtener_logger

INDICES_SQLITE = {
    "listings": ["id", "host_id", "localidad", "categoria_precio"],
    "hosts": ["host_id"],
    "listing_amenities": ["listing_id", "amenity"],
    "calendar": ["listing_id", "date", "anio, mes"],
    "calendar_mensual": ["listing_id", "periodo"],
    "reviews": ["listing_id", "date", "anio, mes"],
}

_CARACTERES_ILEGALES = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class Carga:
    """
    Carga las tablas transformadas en SQLite y XLSX y verifica el resultado.

    Parameters
    ----------
    tablas : dict[str, pandas.DataFrame]
        Salida de ``Transformacion.transformar_todo()``.
    ruta_sqlite : Path
        Archivo de base de datos SQLite a crear.

    Ejemplo
    -------
    >>> carga = Carga(tablas)
    >>> carga.cargar_sqlite()
    >>> carga.exportar_xlsx()
    >>> carga.verificar_carga()
    True
    """

    LOTE_SQLITE = 200_000
    LOTE_EXCEL = 50_000

    def __init__(self, tablas: dict[str, pd.DataFrame], ruta_sqlite: Path = RUTA_SQLITE):
        self.tablas = tablas
        self.ruta_sqlite = Path(ruta_sqlite)
        self.log = obtener_logger("carga")
        self.exportados: dict[Path, list[str]] = {}

    def cargar_sqlite(self) -> None:
        """
        Crea una base de datos SQLite nueva e inserta cada tabla.

        Se elimina el archivo anterior para que cada ejecución genere una base
        "nueva" y reproducible (no se mezclan datos de corridas anteriores).
        Durante la carga masiva se desactiva la sincronización a disco
        (``PRAGMA synchronous=OFF``), lo que acelera la inserción de los
        ~7 millones de filas del calendario; al terminar se crean índices.
        """
        self.log.info("=== Inicio de la fase de CARGA ===")
        self.ruta_sqlite.parent.mkdir(parents=True, exist_ok=True)
        if self.ruta_sqlite.exists():
            self.ruta_sqlite.unlink()
            self.log.info("Base SQLite anterior eliminada: %s", self.ruta_sqlite.name)

        con = sqlite3.connect(self.ruta_sqlite)
        try:
            con.execute("PRAGMA synchronous = OFF")
            con.execute("PRAGMA journal_mode = MEMORY")
            for nombre, df in self.tablas.items():
                inicio = time.perf_counter()
                for i, desde in enumerate(range(0, max(len(df), 1), self.LOTE_SQLITE)):
                    df.iloc[desde:desde + self.LOTE_SQLITE].to_sql(
                        nombre, con, if_exists="replace" if i == 0 else "append", index=False)
                self.log.info("SQLite: tabla '%s' cargada con %s registros (%.1fs)",
                              nombre, f"{len(df):,}", time.perf_counter() - inicio)
            for tabla, columnas in INDICES_SQLITE.items():
                for cols in columnas:
                    nombre_idx = f"idx_{tabla}_{cols.replace(', ', '_')}"
                    con.execute(f"CREATE INDEX IF NOT EXISTS {nombre_idx} ON {tabla} ({cols})")
            con.commit()
            self.log.info("SQLite: índices creados. Base de datos en %s (%.1f MB)",
                          self.ruta_sqlite, self.ruta_sqlite.stat().st_size / 1e6)
        except Exception as e:
            self.log.error("Error cargando datos en SQLite: %s", e, exc_info=True)
            raise
        finally:
            con.close()

    @staticmethod
    def _filas_para_excel(df: pd.DataFrame):
        """
        Entrega las filas como tuplas de valores nativos de Python, por bloques.

        * NaN / NA -> celda vacía (None)
        * se quitan caracteres de control que Excel no acepta en una celda
        """
        for desde in range(0, len(df), Carga.LOTE_EXCEL):
            bloque = df.iloc[desde:desde + Carga.LOTE_EXCEL]
            for col in bloque.select_dtypes(include=["object", "string", "category"]).columns:
                bloque = bloque.assign(**{col: bloque[col].astype("string")
                                          .str.replace(_CARACTERES_ILEGALES, "", regex=True)})
            bloque = bloque.astype(object).where(bloque.notna(), None)
            yield from bloque.itertuples(index=False, name=None)

    def _escribir_libro(self, ruta: Path, nombres: list[str]) -> None:
        """
        Escribe un libro de Excel con una hoja por tabla.

        Se usa xlsxwriter en modo ``constant_memory``: cada fila se escribe en
        disco apenas se genera, así exportar ~500 mil reseñas no requiere
        tener millones de celdas en memoria.
        """
        hojas = []
        libro = xlsxwriter.Workbook(ruta, {
            "constant_memory": True,
            "strings_to_formulas": False, "strings_to_urls": False,
            "nan_inf_to_errors": True,
        })
        negrita = libro.add_format({"bold": True, "bg_color": "#DDEBF7"})
        try:
            for nombre in nombres:
                df = self.tablas[nombre]
                if len(df) > MAX_FILAS_EXCEL:
                    self.log.warning("XLSX: '%s' tiene %s filas y supera el límite de Excel (%s); "
                                     "no se exporta a XLSX (queda completa en SQLite)",
                                     nombre, f"{len(df):,}", f"{MAX_FILAS_EXCEL:,}")
                    continue
                inicio = time.perf_counter()
                hoja = libro.add_worksheet(nombre[:31])
                hoja.write_row(0, 0, list(df.columns), negrita)
                hoja.freeze_panes(1, 0)
                for i, fila in enumerate(self._filas_para_excel(df), start=1):
                    hoja.write_row(i, 0, fila)
                hojas.append(nombre)
                self.log.info("XLSX: hoja '%s' exportada con %s filas (%.1fs)",
                              nombre[:31], f"{len(df):,}", time.perf_counter() - inicio)
        finally:
            libro.close()
        self.exportados[ruta] = hojas
        self.log.info("XLSX: archivo generado %s (%.1f MB)", ruta.name, ruta.stat().st_size / 1e6)

    def exportar_xlsx(self) -> None:
        """
        Exporta las tablas a Excel:

        * ``airbnb_bogota_resumen.xlsx``: todas las tablas que caben en una hoja
          (listings, hosts, amenities, resúmenes de calendario y reseñas,
          auditoría).
        * ``airbnb_bogota_reviews.xlsx``: las reseñas limpias (≈528 mil filas),
          en un libro aparte para que el resumen abra rápido.
        * El calendario diario (≈7 millones de filas) se omite con WARNING.
        """
        try:
            RUTA_XLSX_RESUMEN.parent.mkdir(parents=True, exist_ok=True)
            resumen = [n for n in self.tablas if n != "reviews"]
            self._escribir_libro(RUTA_XLSX_RESUMEN, resumen)
            self._escribir_libro(RUTA_XLSX_REVIEWS, ["reviews"])
        except Exception as e:
            self.log.error("Error exportando a XLSX: %s", e, exc_info=True)
            raise

    def verificar_carga(self) -> bool:
        """
        Comprueba que los datos se hayan cargado correctamente:

        * SQLite: ``SELECT COUNT(*)`` de cada tabla = filas del DataFrame.
        * SQLite: el precio quedó numérico y las fechas en formato YYYY-MM-DD.
        * XLSX: cada hoja tiene las filas esperadas (+1 de encabezado).

        Returns
        -------
        bool
            True si todas las verificaciones pasan.
        """
        self.log.info("--- Verificación de la carga ---")
        ok = True
        con = sqlite3.connect(self.ruta_sqlite)
        try:
            for nombre, df in self.tablas.items():
                n_db = con.execute(f"SELECT COUNT(*) FROM {nombre}").fetchone()[0]
                if n_db == len(df):
                    self.log.info("Verificación SQLite OK  %-26s %s = %s",
                                  nombre, f"{len(df):,}", f"{n_db:,}")
                else:
                    ok = False
                    self.log.error("Verificación SQLite FALLÓ %-26s esperado %s, cargado %s",
                                   nombre, f"{len(df):,}", f"{n_db:,}")

            tipos_precio = con.execute("SELECT DISTINCT typeof(price) FROM listings").fetchall()
            fechas_malas = con.execute(
                "SELECT COUNT(*) FROM calendar WHERE date NOT GLOB "
                "'[0-9][0-9][0-9][0-9]-[0-9][0-9]-[0-9][0-9]'").fetchone()[0]
            if tipos_precio == [("real",)] and fechas_malas == 0:
                self.log.info("Verificación de tipos OK: price es REAL y todas las fechas son YYYY-MM-DD")
            else:
                ok = False
                self.log.error("Verificación de tipos FALLÓ: typeof(price)=%s, fechas inválidas=%d",
                               tipos_precio, fechas_malas)
        finally:
            con.close()

        for ruta, hojas in self.exportados.items():
            libro = load_workbook(ruta, read_only=True)
            for nombre in hojas:
                filas = libro[nombre[:31]].max_row - 1
                if filas == len(self.tablas[nombre]):
                    self.log.info("Verificación XLSX OK    %s / %-20s %s filas",
                                  ruta.name, nombre, f"{filas:,}")
                else:
                    ok = False
                    self.log.error("Verificación XLSX FALLÓ %s / %s: esperado %s, encontrado %s",
                                   ruta.name, nombre, f"{len(self.tablas[nombre]):,}", f"{filas:,}")
            libro.close()

        if ok:
            self.log.info("Carga verificada correctamente: todos los registros coinciden")
        else:
            self.log.error("La verificación de la carga encontró diferencias (ver líneas ERROR)")
        return ok
